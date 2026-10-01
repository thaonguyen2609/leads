
import os
import re
import copy
import unicodedata
from datetime import datetime
from pathlib import Path
import pandas as pd
import numpy as np
import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

# ==========================================
# CẤU HÌNH CHUNG
# ==========================================
FILE_PATH   = r"tool\PO VC26 tuần 03.08.xlsx"
FILE_OUTPUT = r"tool\PO VC26 tuần 03.08_fill.xlsx"
DRY_RUN     = False         

TARGET_SHEETS = None        

# ==========================================
# CẤU HÌNH PHASE 1
# ==========================================
HEADER_KEYWORDS = [
    "cơ cấu", "món ăn", "nguyên liệu", "định lượng",
    "số suất", "đơn giá", "thành tiền", "gọi hàng", "cái hạng", "tổng",
]

SUMMARY_KEYWORDS_BONUS = [
    "cộng", "tổng", "total", "sum", "subtotal",
    "cost id", "giá vị", "gas", "vat", "chiết khấu",
    "phụ thu", "phụ phí", "khay",
]

LAST_CHECK_COLUMN_KEYWORD   = "tổng"
DEFAULT_LAST_COLUMN_LETTER  = "K"

NA_VALUES    = {"n/a", "na", "#n/a", "null", "none", "-", "--", "nil"}
EXCEL_ERRORS = {"#n/a", "#ref!", "#div/0!", "#value!", "#name?", "#num!", "#null!"}

MAX_ISSUES_PRINT = 50

# ==========================================
# CẤU HÌNH PHASE 2
# ==========================================
SHEETS_CAN_XU_LY = {
    "thu 2", "thu 3", "thu 4", "thu 5", "thu 6", "thu 7", "chu nhat"
}

HEADER_CHUAN = [
    "Cơ cấu", "Món ăn", "Mã hàng", "Nguyên liệu", "Định lượng sống",
    "Số suất", "Gọi hàng", "Đơn giá", "Thành tiền", "Tổng",
]

HEADER_MAP = {
    0: ("co cau",),
    1: ("mon an",),
    2: ("ma hang", "vinalpha", "ma ncc", "ma sp", "ma mon"),
    3: ("nguyen lieu",),
    4: ("dinh luong",),
    5: ("so suat",),
    6: ("goi hang",),
    7: ("don gia",),
    8: ("thanh tien",),
    9: ("tong",),
}
NGUONG_KHOP = 4

# ==========================================
# CẤU HÌNH PHASE 3 / 4
# ==========================================
SHEET_BAOGIA  = "Báo giá"
SHEET_TARGETS = ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ nhật"]

COL_MA_HANG      = "Mã hàng"
COL_TEN_MAT_HANG = "Tên mặt hàng"
COL_DON_GIA      = "Đơn giá"
COL_NGUYEN_LIEU  = "Nguyên liệu"


# ==========================================
# HÀM HỖ TRỢ CHUNG
# ==========================================

def chuan_hoa(s):
    """Chuẩn hoá chuỗi tiếng Việt (bỏ dấu, viết thường, thay 'đ'→'d')."""
    if s is None:
        return ""
    s = re.sub(r"\s+", " ", str(s)).strip().lower()
    s = s.replace("đ", "d")
    return "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )


def parse_number(val):
    if val is None: return None
    if isinstance(val, float) and pd.isna(val): return None
    if isinstance(val, (int, float)): return float(val)
    s = str(val).replace(",", "").replace(".", "").replace("%", "").strip()
    if s.replace("-", "").isdigit():
        try: return float(s)
        except: return None
    return None


def is_empty(val):
    if val is None: return True
    if isinstance(val, float) and pd.isna(val): return True
    if isinstance(val, str) and val.strip() == "": return True
    return False


def is_na_or_error(val):
    if val is None: return None, None
    if isinstance(val, float) and pd.isna(val): return None, None
    s = str(val).strip().lower()
    if s in NA_VALUES:    return "NA", val
    if s in EXCEL_ERRORS: return "ERROR", val
    return None, None


def row_to_text(row):
    vals = [str(v).lower() for v in row if not is_empty(v)]
    return " | ".join(vals)


def is_header_row(row):
    text = row_to_text(row)
    return sum(1 for kw in HEADER_KEYWORDS if kw in text) >= 3


def is_title_row(row):
    vals = [v for v in row if not is_empty(v)]
    if not (1 <= len(vals) <= 3): return False
    for v in vals:
        if isinstance(v, (int, float)): return False
        if isinstance(v, str) and parse_number(v) is not None: return False
    return len(" ".join(str(v) for v in vals)) > 5


def find_last_check_column(header_row, n_cols):
    for c_idx in range(n_cols):
        val = header_row[c_idx] if c_idx < len(header_row) else None
        if is_empty(val): continue
        if LAST_CHECK_COLUMN_KEYWORD in str(val).strip().lower():
            return c_idx
    try:
        return min(column_index_from_string(DEFAULT_LAST_COLUMN_LETTER) - 1, n_cols - 1)
    except:
        return min(10, n_cols - 1)


# ==========================================
# PHASE 1 – LOGIC NHẬN DIỆN DÒNG TỔNG KẾT
# ==========================================

def compute_data_profile(data_rows, last_check_col):
    if not data_rows:
        return {"max_filled": 0, "typical_density": 0}
    filled_counts = []
    for row in data_rows:
        cnt = sum(1 for c in range(last_check_col + 1) if not is_empty(row[c]))
        filled_counts.append(cnt)
    filled_counts.sort()
    core = filled_counts[1:-1] if len(filled_counts) >= 4 else filled_counts
    return {
        "max_filled": max(filled_counts),
        "typical_density": sum(core) / len(core) if core else 0,
    }


def is_summary_row_structural(row, profile, last_check_col):
    nonempty = [v for v in row if not is_empty(v)]
    if len(nonempty) == 0:
        return False
    if is_header_row(row):
        return False

    text_lower = row_to_text(row)
    if any(kw in text_lower for kw in SUMMARY_KEYWORDS_BONUS):
        if len(nonempty) <= 3:
            return True

    filled = sum(1 for c in range(last_check_col + 1) if not is_empty(row[c]))
    crit1 = False
    if profile["max_filled"] > 0:
        crit1 = filled < profile["max_filled"] * 0.5
    if profile["typical_density"] > 0:
        crit1 = crit1 or filled < profile["typical_density"] * 0.6

    middle_start = 2
    middle_end = max(last_check_col - 1, middle_start)
    middle_filled = sum(
        1 for c in range(middle_start, middle_end + 1)
        if c < len(row) and not is_empty(row[c])
    )
    crit2 = (middle_filled == 0)

    num_positions = [
        i for i, v in enumerate(row)
        if i <= last_check_col and not is_empty(v) and parse_number(v) is not None
    ]
    crit3 = False
    if 0 < len(num_positions) <= 2:
        if all(p >= last_check_col - 1 for p in num_positions):
            crit3 = True
    if len(num_positions) == 0 and len(nonempty) <= 2:
        crit3 = True

    return sum([crit1, crit2, crit3]) >= 2


# ==========================================
# ĐỌC DỮ LIỆU (dùng chung cho Phase 1)
# ==========================================

def load_sheet_data(file_path, sheet_name):
    wb_f = openpyxl.load_workbook(file_path, data_only=False)
    wb_v = openpyxl.load_workbook(file_path, data_only=True)
    ws_f = wb_f[sheet_name]
    ws_v = wb_v[sheet_name]

    n_rows, n_cols = ws_f.max_row, ws_f.max_column
    data = [[None] * n_cols for _ in range(n_rows)]
    formulas = {}

    for r in range(1, n_rows + 1):
        for c in range(1, n_cols + 1):
            cf = ws_f.cell(row=r, column=c)
            cv = ws_v.cell(row=r, column=c)
            if isinstance(cf.value, str) and cf.value.startswith("="):
                formulas[(r, c)] = (cf.value, cv.value)
                data[r-1][c-1] = cv.value
            else:
                data[r-1][c-1] = cf.value

    merges = []
    for mr in ws_f.merged_cells.ranges:
        merges.append(str(mr))
        min_c, min_r, max_c, max_r = mr.min_col, mr.min_row, mr.max_col, mr.max_row
        anchor = data[min_r-1][min_c-1]
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                data[r-1][c-1] = anchor

    return pd.DataFrame(data), formulas, merges, wb_f


# ==========================================
# PHASE 1 – PHÂN TÍCH SHEET
# ==========================================

def analyze_sheet(df_raw, formulas, sheet_name):
    n_rows, n_cols = df_raw.shape
    tables = []
    i = 0
    while i < n_rows:
        row = df_raw.iloc[i].tolist()
        if is_header_row(row):
            title, title_row = None, None
            for back in range(1, 4):
                if i - back < 0: break
                if is_title_row(df_raw.iloc[i - back].tolist()):
                    title = " ".join(
                        str(v) for v in df_raw.iloc[i - back].tolist()
                        if not is_empty(v)
                    ).strip()
                    title_row = i - back
                    break

            last_col = find_last_check_column(row, n_cols)
            last_col_letter = get_column_letter(last_col + 1)

            candidate_rows = []
            j = i + 1
            while j < n_rows:
                r = df_raw.iloc[j].tolist()
                nonempty = [v for v in r if not is_empty(v)]
                if len(nonempty) == 0: break
                if is_header_row(r): break
                num_count = sum(1 for v in nonempty if parse_number(v) is not None)
                if num_count == 0 and len(nonempty) <= 3:
                    break
                candidate_rows.append(j)
                j += 1

            candidate_values = [df_raw.iloc[r].tolist() for r in candidate_rows]
            profile = compute_data_profile(candidate_values, last_col)

            data_rows, summary_rows = [], []
            for r_idx in candidate_rows:
                r = df_raw.iloc[r_idx].tolist()
                if is_summary_row_structural(r, profile, last_col):
                    summary_rows.append(r_idx)
                else:
                    data_rows.append(r_idx)

            tables.append({
                "title": title or "(không rõ)",
                "title_row": title_row,
                "header_row": i,
                "data_rows": data_rows,
                "summary_rows": summary_rows,
                "end_row": j - 1,
                "last_check_col": last_col,
                "last_check_col_letter": last_col_letter,
                "profile": profile,
            })
            i = j
        else:
            i += 1

    issues = {
        "missing": [], "zero": [], "na": [], "error": [],
        "formula_none": [], "formula_zero": [],
    }

    for t in tables:
        last_col = t["last_check_col"]
        for r_idx in t["data_rows"]:
            row = df_raw.iloc[r_idx].tolist()
            excel_row = r_idx + 1
            for c_idx in range(last_col + 1):
                val = row[c_idx]
                col_letter = get_column_letter(c_idx + 1)
                key = (excel_row, c_idx + 1)

                if key in formulas:
                    f_str, cached = formulas[key]
                    if cached is None:
                        issues["formula_none"].append((excel_row, col_letter, f_str)); continue
                    kind, _ = is_na_or_error(cached)
                    if kind == "NA":
                        issues["na"].append((excel_row, col_letter, cached, f_str)); continue
                    if kind == "ERROR":
                        issues["error"].append((excel_row, col_letter, cached, f_str)); continue
                    if parse_number(cached) == 0:
                        issues["formula_zero"].append((excel_row, col_letter, cached))
                    continue

                if is_empty(val):
                    issues["missing"].append((excel_row, col_letter)); continue
                kind, _ = is_na_or_error(val)
                if kind == "NA":
                    issues["na"].append((excel_row, col_letter, val, None)); continue
                if kind == "ERROR":
                    issues["error"].append((excel_row, col_letter, val, None)); continue
                if parse_number(val) == 0:
                    issues["zero"].append((excel_row, col_letter))

    return {"n_rows": n_rows, "n_cols": n_cols, "tables": tables, "issues": issues}


def extract_row_classification(info):
    """
    Từ kết quả analyze_sheet, trả về 2 set SỐ DÒNG EXCEL (1-based):
      - data_rows_set    : các dòng DỮ LIỆU
      - summary_rows_set : các dòng TỔNG KẾT
    """
    data_rows_set = set()
    summary_rows_set = set()
    for t in info.get("tables", []):
        for r_idx in t["data_rows"]:
            data_rows_set.add(r_idx + 1)
        for r_idx in t["summary_rows"]:
            summary_rows_set.add(r_idx + 1)
    return data_rows_set, summary_rows_set


# ==========================================
# PHASE 1 – IN BÁO CÁO
# ==========================================

def print_phase1_report(sheet_name, info, merges):
    print("\n" + "="*78)
    print(f"  📄 SHEET: {sheet_name}")
    print("="*78)
    print(f"  Kích thước: {info['n_rows']} dòng × {info['n_cols']} cột")
    print(f"  Ô gộp (merged cells): {len(merges)} vùng")

    print(f"\n  📋 SỐ BẢNG PHÁT HIỆN: {len(info['tables'])}")
    for k, t in enumerate(info["tables"], 1):
        print(f"\n  ── Bảng {k}: {t['title']}")
        if t['title_row'] is not None:
            print(f"     • Dòng tiêu đề: {t['title_row'] + 1}")
        print(f"     • Dòng header : {t['header_row'] + 1}")
        print(f"     • Số dòng DỮ LIỆU  : {len(t['data_rows'])}")
        print(f"     • Số dòng TỔNG KẾT : {len(t['summary_rows'])}")
        if t['summary_rows']:
            srows = ", ".join(str(r + 1) for r in t['summary_rows'])
            print(f"        → Dòng bị coi là tổng kết: {srows}")
        print(f"     • Profile: max_filled={t['profile']['max_filled']}, "
              f"typical_density={t['profile']['typical_density']:.1f}")
        print(f"     • Vùng kiểm tra: cột A → cột {t['last_check_col_letter']}")

    iss = info["issues"]
    total = (len(iss["missing"]) + len(iss["zero"]) +
             len(iss["na"]) + len(iss["error"]))
    print(f"\n  ⚠️  VẤN ĐỀ DỮ LIỆU: {total} ô")

    if iss["missing"]:
        print(f"\n     [KHUYẾT] – {len(iss['missing'])} ô")
        for r, c in iss["missing"][:MAX_ISSUES_PRINT]:
            print(f"        · Dòng {r}, cột {c}")
        if len(iss["missing"]) > MAX_ISSUES_PRINT:
            print(f"        ... và {len(iss['missing']) - MAX_ISSUES_PRINT} ô khác")
    if iss["zero"]:
        print(f"\n     [BẰNG 0] – {len(iss['zero'])} ô")
        for r, c in iss["zero"][:MAX_ISSUES_PRINT]:
            print(f"        · Dòng {r}, cột {c}")
        if len(iss["zero"]) > MAX_ISSUES_PRINT:
            print(f"        ... và {len(iss['zero']) - MAX_ISSUES_PRINT} ô khác")
    if iss["na"]:
        print(f"\n     [N/A] – {len(iss['na'])} ô")
        for r, c, val, f in iss["na"][:MAX_ISSUES_PRINT]:
            note = f"  ← công thức: {f}" if f else "  ← gõ tay"
            print(f"        · Dòng {r}, cột {c}: {val}{note}")
    if iss["error"]:
        print(f"\n     [LỖI CÔNG THỨC] – {len(iss['error'])} ô")
        for r, c, val, f in iss["error"][:MAX_ISSUES_PRINT]:
            note = f"  ← công thức: {f}" if f else "  ← gõ tay"
            print(f"        · Dòng {r}, cột {c}: {val}{note}")
    if iss["formula_none"]:
        print(f"\n     [CÔNG THỨC CHƯA TÍNH] – {len(iss['formula_none'])} ô")
        for r, c, f in iss["formula_none"][:MAX_ISSUES_PRINT]:
            print(f"        · Dòng {r}, cột {c}: {f}")
    if iss["formula_zero"]:
        print(f"\n     [CÔNG THỨC = 0] – {len(iss['formula_zero'])} ô")
        for r, c, v in iss["formula_zero"][:MAX_ISSUES_PRINT]:
            print(f"        · Dòng {r}, cột {c} → giá trị = {v}")


# ==========================================
# PHASE 2 – KIỂM TRA HEADER
# ==========================================

def kiem_tra_header_dong(row):
    """Trả về (là_header, ds_loi). ds_loi = [(vị_trí, giá_trị_cũ), ...]."""
    so_khop = 0
    for i in range(10):
        if i >= len(row): break
        if any(kw in chuan_hoa(row[i].value) for kw in HEADER_MAP[i]):
            so_khop += 1
    if so_khop < NGUONG_KHOP:
        return False, []

    ds_loi = []
    for i in range(10):
        if i >= len(row): break
        if row[i].value != HEADER_CHUAN[i]:
            ds_loi.append((i, row[i].value))
    return True, ds_loi


def run_phase2(wb, sheet_name):
    """Trả về (so_header, so_loi, chi_tiet) cho sheet."""
    if chuan_hoa(sheet_name) not in SHEETS_CAN_XU_LY:
        return 0, 0, []

    ws = wb[sheet_name]
    so_header = 0
    so_loi = 0
    chi_tiet = []

    for row in ws.iter_rows():
        ok, ds_loi = kiem_tra_header_dong(row)
        if not ok:
            continue
        so_header += 1
        if not ds_loi:
            continue
        so_loi += 1
        r = row[0].row
        chi_tiet.append((r, ds_loi))

        if not DRY_RUN:
            for i, ten in enumerate(HEADER_CHUAN):
                if i < len(row):
                    row[i].value = ten

    return so_header, so_loi, chi_tiet


# ==========================================
# PHASE 3 – VLOOKUP RESOLVER
# ==========================================

VLOOKUP_RE = re.compile(
    r"""=\s*VLOOKUP\s*\(\s*
        \$?([A-Z]+)\$?(\d+)\s*,\s*
        '?([^'!]+?)'?\s*!\s*
        \$?([A-Z]+)\$?(\d+)?\s*:\s*
        \$?([A-Z]+)\$?(\d+)?\s*,\s*
        (\d+)\s*,\s*
        (0|1|FALSE|TRUE|false|true)\s*\)
    """,
    re.VERBOSE | re.IGNORECASE
)


def _norm_key(v):
    if v is None:
        return ""
    if isinstance(v, float) and np.isnan(v):
        return ""
    s = str(v).strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _resolve_one_vlookup(formula, ws_current, wb_read, baogia_sheet):
    m = VLOOKUP_RE.match(formula.strip())
    if not m:
        return None

    key_col = m.group(1)
    key_row = int(m.group(2))
    source_sheet_name = m.group(3).strip().strip("'")
    start_col = m.group(4)
    start_row = int(m.group(5)) if m.group(5) else 1
    end_col = m.group(6)
    end_row = int(m.group(7)) if m.group(7) else None
    result_col_idx = int(m.group(8))
    match_type_str = m.group(9).upper()
    match_type = 0 if match_type_str in ("0", "FALSE") else 1

    source_sheet = None
    for s in wb_read.sheetnames:
        if s.strip().lower() == source_sheet_name.lower():
            source_sheet = s
            break
    if source_sheet is None:
        return None

    ws_src = wb_read[source_sheet]

    key_coord = f"{key_col}{key_row}"
    try:
        key_value = ws_current[key_coord].value
    except Exception:
        return None

    if isinstance(key_value, str) and key_value.startswith("="):
        return None

    key_norm = _norm_key(key_value)
    if not key_norm:
        return None

    try:
        start_c = column_index_from_string(start_col)
        end_c = column_index_from_string(end_col)
    except Exception:
        return None

    if end_row is None:
        end_row = ws_src.max_row

    result_c = start_c + result_col_idx - 1
    if result_c > end_c:
        return None

    for r in range(start_row, end_row + 1):
        cell_val = ws_src.cell(row=r, column=start_c).value
        if cell_val is None:
            continue
        if match_type == 0:
            if _norm_key(cell_val) == key_norm:
                result = ws_src.cell(row=r, column=result_c).value
                if result is not None:
                    return result
                return None

    return None


def _detect_header_cells(ws):
    header_row = None
    headers = {}
    for r in range(1, min(16, ws.max_row + 1)):
        temp = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                temp[str(v).strip()] = c
        if COL_NGUYEN_LIEU in temp and COL_DON_GIA in temp and COL_MA_HANG in temp:
            header_row = r
            headers = temp
            break

    if header_row is None:
        return None, {}, {}, set()

    header_row_values = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        header_row_values[c] = str(v).strip() if v is not None else ""

    header_cells = set()
    for c in range(1, ws.max_column + 1):
        header_cells.add((header_row, c))

    return header_row, headers, header_row_values, header_cells


def resolve_vlookups_in_sheet(sheet_name, wb_write, wb_read, wb_formula, baogia_sheet,
                              data_rows_set=None):
    """
    Resolve VLOOKUP. Nếu có data_rows_set (từ Phase 1) → CHỈ resolve
    trên các dòng đó.
    """
    if sheet_name not in wb_write.sheetnames:
        return 0

    ws_w = wb_write[sheet_name]
    ws_v = wb_read[sheet_name]
    ws_f = wb_formula[sheet_name]

    _, _, _, header_cells = _detect_header_cells(ws_w)

    total_resolved = 0
    for _ in range(3):
        changed = 0
        for row in ws_f.iter_rows():
            for cell in row:
                if (cell.row, cell.column) in header_cells:
                    continue

                if data_rows_set is not None and cell.row not in data_rows_set:
                    continue

                v = cell.value
                if not (isinstance(v, str) and v.lstrip().upper().startswith("=VLOOKUP")):
                    continue

                result = _resolve_one_vlookup(v, ws_v, wb_read, baogia_sheet)

                if result is not None and result != "":
                    if ws_w is not ws_f:
                        ws_w[cell.coordinate].value = result
                    ws_v[cell.coordinate].value = result
                    ws_f[cell.coordinate].value = result
                    changed += 1
                    total_resolved += 1

        if changed == 0:
            break

    return total_resolved


# ==========================================
# PHASE 3 – HÀM HỖ TRỢ
# ==========================================

def is_blank(val):
    if val is None:
        return True
    if isinstance(val, float) and np.isnan(val):
        return True
    s = str(val).strip()
    return s == "" or s.lower() in ("nan", "none", "#n/a")


def is_header_like_row(ws, row_idx, header_row_values, min_match=2):
    matched = 0
    for c, hdr_val in header_row_values.items():
        if not hdr_val:
            continue
        v = ws.cell(row=row_idx, column=c).value
        if v is None:
            continue
        if str(v).strip() == hdr_val:
            matched += 1
    return matched >= min_match


def normalize_price(val):
    if is_blank(val):
        return None
    if isinstance(val, (int, float, np.integer, np.floating)):
        try:
            return float(val)
        except Exception:
            return None
    s = str(val).strip()
    for token in ["VNĐ", "VND", "vnđ", "vnd", "VNd", "đ", "Đ", "$", " ", "\u00a0"]:
        s = s.replace(token, "")
    has_dot = "." in s
    has_comma = "," in s
    try:
        if has_dot and has_comma:
            if s.rfind(".") > s.rfind(","):
                s = s.replace(",", "")
            else:
                s = s.replace(".", "").replace(",", ".")
        elif has_dot:
            parts = s.split(".")
            if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
                s = s.replace(".", "")
        elif has_comma:
            parts = s.split(",")
            if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
                s = s.replace(",", "")
            else:
                s = s.replace(",", ".")
        return float(s)
    except Exception:
        return None


def norm_text(val):
    if is_blank(val):
        return ""
    s = str(val).strip().lower()
    return re.sub(r"\s+", " ", s)


def find_sheet(sheet_list, target_name):
    for s in sheet_list:
        if s.strip().lower() == target_name.strip().lower():
            return s
    return None


def read_baogia_auto(input_path, sheet_name, required_cols, max_scan_rows=15):
    df_raw = pd.read_excel(input_path, sheet_name=sheet_name, header=None)

    header_row = None
    for r in range(min(max_scan_rows, len(df_raw))):
        row_vals = [str(v).strip() if v is not None else "" for v in df_raw.iloc[r].tolist()]
        if all(any(rc == v for v in row_vals) for rc in required_cols):
            header_row = r
            break

    if header_row is None:
        print(f"⚠️ Không dò được header trong {max_scan_rows} dòng đầu của '{sheet_name}'.")
        return None

    print(f"📌 Header thực ở dòng {header_row + 1} của '{sheet_name}'")
    df = pd.read_excel(input_path, sheet_name=sheet_name, header=header_row)
    df.columns = [str(c).strip() for c in df.columns]

    keep_cols = []
    for c in df.columns:
        if c.startswith("Unnamed") and df[c].notna().sum() == 0:
            continue
        keep_cols.append(c)
    df = df[keep_cols]
    print(f"   Các cột nhận diện: {list(df.columns)}")
    return df


# ==========================================
# PHASE 3 – XỬ LÝ 1 SHEET
# ==========================================

def process_sheet(wb_write, wb_read, wb_read_formula, sheet_name,
                  lookup_exact, lookup_name,
                  data_rows_set=None, summary_rows_set=None):
    """
    data_rows_set / summary_rows_set: tập SỐ DÒNG EXCEL (1-based) do Phase 1
    phân loại. Nếu có, Phase 3 chỉ duyệt trên data_rows_set.
    """
    print("\n" + "=" * 65)
    print(f"🔧 XỬ LÝ SHEET: {sheet_name}")
    print("=" * 65)

    ws_w = wb_write[sheet_name]
    ws_r = wb_read[sheet_name]
    ws_f = wb_read_formula[sheet_name]

    header_row, headers, header_row_values, _ = _detect_header_cells(ws_w)

    if header_row is None:
        print(f"  ⚠️ Không tìm thấy cột {[COL_NGUYEN_LIEU, COL_DON_GIA, COL_MA_HANG]}. Bỏ qua.")
        return

    col_nl = headers[COL_NGUYEN_LIEU]
    col_dg = headers[COL_DON_GIA]
    col_mh = headers[COL_MA_HANG]
    print(f"  📌 Header ở dòng {header_row}")
    print(f"  📌 Cột: NL={col_nl}, ĐG={col_dg}, MH={col_mh}")

    protected_names = {
        norm_text(COL_NGUYEN_LIEU), norm_text(COL_DON_GIA),
        norm_text(COL_MA_HANG), norm_text(COL_TEN_MAT_HANG),
    }

    if data_rows_set is not None:
        iterate_rows = sorted(data_rows_set)
        print(f"  🎯 Phase 1 cung cấp {len(iterate_rows)} dòng DỮ LIỆU "
              f"(bỏ qua {len(summary_rows_set or set())} dòng TỔNG KẾT)")
    else:
        iterate_rows = list(range(header_row + 1, ws_w.max_row + 1))
        print(f"  ⚠️ Không có classification từ Phase 1 → duyệt toàn bộ "
              f"{len(iterate_rows)} dòng từ sau header")

    # ---- BƯỚC 1 ----
    print(f"\n  🔹 BƯỚC 1: Khớp chính xác (Nguyên liệu + Đơn giá)")
    to_buoc2 = []
    full_data_no_match = []
    matched = 0
    blank_price_with_formula = 0
    blank_price_no_formula = 0
    skipped_header_rows = 0

    for row_idx in iterate_rows:

        if is_header_like_row(ws_w, row_idx, header_row_values, min_match=2):
            skipped_header_rows += 1
            continue

        nl = ws_r.cell(row=row_idx, column=col_nl).value
        dg = ws_r.cell(row=row_idx, column=col_dg).value

        if is_blank(nl):
            continue

        if norm_text(nl) in protected_names:
            skipped_header_rows += 1
            continue

        if is_blank(dg):
            f_val = ws_f.cell(row=row_idx, column=col_dg).value
            has_formula = isinstance(f_val, str) and f_val.startswith("=")
            if has_formula:
                blank_price_with_formula += 1
                to_buoc2.append((row_idx, nl, dg, "công thức chưa resolve"))
            else:
                blank_price_no_formula += 1
                to_buoc2.append((row_idx, nl, dg, "đơn giá trống"))
            continue

        ten = norm_text(nl)
        price = normalize_price(dg)
        key = (ten, price)
        if key in lookup_exact:
            ws_w.cell(row=row_idx, column=col_mh).value = lookup_exact[key]
            matched += 1
        else:
            ws_w.cell(row=row_idx, column=col_mh).value = None
            full_data_no_match.append((row_idx, nl, dg))

    print(f"     ✓ Khớp chính xác             : {matched} dòng")
    print(f"     ⊘ Đủ NL+ĐG nhưng không khớp  : {len(full_data_no_match)} dòng → ĐỂ TRỐNG")
    print(f"     ↳ Chuyển sang Bước 2 (thiếu ĐG): {len(to_buoc2)} dòng")
    if blank_price_with_formula:
        print(f"        • Do công thức chưa resolve : {blank_price_with_formula}")
    if blank_price_no_formula:
        print(f"        • Do đơn giá trống thật     : {blank_price_no_formula}")
    if skipped_header_rows:
        print(f"     ⊘ Bỏ qua {skipped_header_rows} dòng giống header lặp lại")

    if full_data_no_match:
        print(f"\n     📌 Danh sách dòng ĐỦ NL+ĐG nhưng KHÔNG có trong Báo giá (đã để trống mã):")
        for r, nl, dg in full_data_no_match:
            print(f"        - Dòng {r}: NL='{nl}' | ĐG='{dg}'")

    # ---- BƯỚC 2 ----
    print(f"\n  🔹 BƯỚC 2: Xử lý {len(to_buoc2)} dòng thiếu Đơn giá")
    if to_buoc2:
        print("     📌 Danh sách dòng thiếu ĐG:")
        for r, nl, dg, reason in to_buoc2:
            print(f"        - Dòng {r}: NL='{nl}' | ĐG='{dg}' | {reason}")

    updated = 0
    cleared = 0
    for row_idx, nl, dg, reason in to_buoc2:
        ten = norm_text(nl)
        if ten in lookup_name:
            candidates = lookup_name[ten]
            min_price, best_ma = min(candidates, key=lambda x: x[0])
            ws_w.cell(row=row_idx, column=col_mh).value = best_ma
            ws_w.cell(row=row_idx, column=col_dg).value = min_price
            updated += 1
            print(f"        ✓ Dòng {row_idx}: Gán '{best_ma}' | Giá thấp nhất = {min_price}")
        else:
            ws_w.cell(row=row_idx, column=col_mh).value = None
            cleared += 1
            print(f"        ✗ Dòng {row_idx}: Không có NL '{nl}' trong Báo giá → trống")

    print(f"     → Cập nhật giá thấp nhất : {updated}")
    print(f"     → Để trống mã hàng       : {cleared}")

    # ---- BƯỚC 3 ----
    print(f"\n  🔹 BƯỚC 3: Dòng có NL trống nhưng ≥2 ô khác có giá trị")
    warnings = []

    check_rows = set()
    if data_rows_set is not None:
        check_rows |= data_rows_set
    if summary_rows_set is not None:
        check_rows |= summary_rows_set
    if not check_rows:
        check_rows = set(range(header_row + 1, ws_w.max_row + 1))

    for row_idx in sorted(check_rows):

        if is_header_like_row(ws_w, row_idx, header_row_values, min_match=2):
            continue

        nl = ws_r.cell(row=row_idx, column=col_nl).value
        if not is_blank(nl):
            continue

        filled = 0
        info = {}
        for name, c in headers.items():
            if c == col_nl:
                continue
            v = ws_r.cell(row=row_idx, column=c).value
            if is_blank(v):
                fv = ws_f.cell(row=row_idx, column=c).value
                if isinstance(fv, str) and fv.startswith("="):
                    v = f"[CT] {fv}"
            if not is_blank(v):
                filled += 1
                info[name] = v

        if filled >= 2:
            if summary_rows_set and row_idx in summary_rows_set:
                classify = "SUMMARY (Phase 1 bỏ qua)"
            elif data_rows_set and row_idx in data_rows_set:
                classify = "DATA (bất thường!)"
            else:
                classify = "OTHER"
            warnings.append((row_idx, info, classify))

    if warnings:
        print(f"     ⚠️ Có {len(warnings)} dòng cần kiểm tra:")
        for r, info, classify in warnings:
            print(f"        - Dòng {r} [{classify}]: {info}")
    else:
        print("     ✓ Không có dòng bất thường.")


# ==========================================
# PHASE 4 – XÓA DÒNG DATA CÓ [NGUYÊN LIỆU] TRỐNG
# ==========================================

# Regex khớp cell reference: [$]?[A-Z]{1,3}[$]?[0-9]+
CELL_REF_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def _shift_formula_rows(formula, deleted_rows_sorted):
    """
    Điều chỉnh row references trong formula sau khi xóa dòng.
      • Ref tới dòng đã bị xóa  → #REF!
      • Ref tới dòng còn sống    → trừ đi số dòng đã xóa phía trên nó
    Hỗ trợ: A5, $A$5, A$5, $A5, range A5:B10, SUM(A5:A10), v.v.
    """
    deleted_set = set(deleted_rows_sorted)

    def repl(m):
        dollar_col, col, dollar_row, row_str = m.groups()
        row_num = int(row_str)
        if row_num in deleted_set:
            return f"{dollar_col}{col}{dollar_row}#REF!"
        delta = sum(1 for r in deleted_rows_sorted if r < row_num)
        return f"{dollar_col}{col}{dollar_row}{row_num - delta}"

    return CELL_REF_RE.sub(repl, formula)


def adjust_formulas_after_delete(ws, deleted_rows):
    """
    openpyxl.delete_rows() KHÔNG tự điều chỉnh formula refs → ta tự làm.
    """
    deleted_sorted = sorted(deleted_rows)
    n_adjusted = 0
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            try:
                new_v = _shift_formula_rows(v, deleted_sorted)
                if new_v != v:
                    cell.value = new_v
                    n_adjusted += 1
            except Exception as e:
                print(f"  ⚠️ Không adjust được formula tại {cell.coordinate}: {e}")
    return n_adjusted


def delete_blank_nl_data_rows(wb, sheet_name, data_rows_set, headers):
    """
    Xóa các dòng DATA có ô [Nguyên liệu] trống — AN TOÀN:
      • Chỉ xóa dòng trong data_rows_set (không đụng header/title/summary).
      • Backup merges + ANCHOR VALUE; nếu anchor row bị xóa → chuyển value
        sang dòng sống kế tiếp trong cùng vùng merge (giữ hiển thị "Món ăn").
      • Sau khi xóa, TỰ ĐIỀU CHỈNH row refs của mọi formula.
    Trả về: (rows_deleted, n_merge_ok, n_merge_skip, n_formula_adjusted)
    """
    ws = wb[sheet_name]
    col_nl = headers[COL_NGUYEN_LIEU]

    # ----- 1. Xác định dòng cần xóa -----
    rows_to_delete = []
    for r in sorted(data_rows_set):
        if r > ws.max_row:
            continue
        v = ws.cell(row=r, column=col_nl).value
        if is_blank(v):
            rows_to_delete.append(r)

    if not rows_to_delete:
        return [], 0, 0, 0

    deleted_set = set(rows_to_delete)

    # ----- 2. Backup merges + anchor value, rồi unmerge -----
    merges_backup = []
    for mr in list(ws.merged_cells.ranges):
        min_r, min_c, max_r, max_c = mr.min_row, mr.min_col, mr.max_row, mr.max_col
        anchor_value = ws.cell(row=min_r, column=min_c).value
        merges_backup.append({
            "min_r": min_r, "min_c": min_c,
            "max_r": max_r, "max_c": max_c,
            "anchor_value": anchor_value,
        })
        ws.unmerge_cells(str(mr))

    # ----- 3. Nếu anchor row bị xóa → chuyển value sang dòng sống kế tiếp -----
    n_anchor_moved = 0
    for m in merges_backup:
        if m["min_r"] in deleted_set and m["anchor_value"] is not None:
            surviving = [
                r for r in range(m["min_r"] + 1, m["max_r"] + 1)
                if r not in deleted_set
            ]
            if surviving:
                new_anchor_r = surviving[0]
                ws.cell(row=new_anchor_r, column=m["min_c"]).value = m["anchor_value"]
                n_anchor_moved += 1

    # ----- 4. Xóa dòng từ dưới lên -----
    for r in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(r, 1)

    # ----- 5. Re-merge với tọa độ đã dịch -----
    n_ok = 0
    n_skip = 0
    for m in merges_backup:
        before = sum(1 for r in rows_to_delete if r < m["min_r"])
        inside = sum(1 for r in rows_to_delete if m["min_r"] <= r <= m["max_r"])

        if inside == (m["max_r"] - m["min_r"] + 1):
            n_skip += 1
            continue

        new_min_r = m["min_r"] - before
        new_max_r = m["max_r"] - before - inside

        if new_max_r < new_min_r:
            n_skip += 1
            continue

        try:
            ws.merge_cells(
                start_row=new_min_r, start_column=m["min_c"],
                end_row=new_max_r,   end_column=m["max_c"]
            )
            n_ok += 1
        except Exception as e:
            print(f"  ⚠️ Re-merge lỗi ({m['min_r']},{m['min_c']})-"
                  f"({m['max_r']},{m['max_c']}): {e}")
            n_skip += 1

    # ----- 6. Điều chỉnh công thức -----
    n_formula = adjust_formulas_after_delete(ws, rows_to_delete)

    if n_anchor_moved:
        print(f"     • Đã chuyển anchor của {n_anchor_moved} vùng merge")

    return rows_to_delete, n_ok, n_skip, n_formula


# ==========================================
# MAIN
# ==========================================

def main():
    if not os.path.exists(FILE_PATH):
        print(f"[LỖI] Không tìm thấy file: {FILE_PATH}")
        return

    print("╔" + "═"*76 + "╗")
    print("║  TOOL KIỂM TRA + ĐIỀN MÃ HÀNG EXCEL DỰ TOÁN (TÍCH HỢP)                 ║")
    print("╚" + "═"*76 + "╝")
    print(f"  File input : {FILE_PATH}")
    print(f"  File output: {FILE_OUTPUT}")
    print(f"  Chế độ     : {'DRY_RUN (không ghi file)' if DRY_RUN else 'GHI FILE OUTPUT'}")

    xls = pd.ExcelFile(FILE_PATH)
    sheets = TARGET_SHEETS if TARGET_SHEETS else xls.sheet_names

    grand = {
        "tables": 0, "data_rows": 0, "summary_rows": 0,
        "missing": 0, "zero": 0, "na": 0, "error": 0,
        "fzero": 0, "fnone": 0,
        "headers": 0, "header_errors": 0,
        "deleted_rows": 0, "deleted_merges_ok": 0, "deleted_merges_skip": 0,
        "adjusted_formulas": 0,
    }

    sheet_analysis_map = {}
    sheet_classify_map = {}

    # =====================================================
    # PHASE 1
    # =====================================================
    print("\n\n" + "█"*78)
    print("█  PHASE 1 – BÁO CÁO CẤU TRÚC & LỖI DỮ LIỆU")
    print("█"*78)

    for sheet in sheets:
        try:
            df_raw, formulas, merges, _ = load_sheet_data(FILE_PATH, sheet)
            info = analyze_sheet(df_raw, formulas, sheet)
            print_phase1_report(sheet, info, merges)

            sheet_analysis_map[sheet] = info
            sheet_classify_map[sheet] = extract_row_classification(info)

            grand["tables"]       += len(info["tables"])
            grand["data_rows"]    += sum(len(t["data_rows"])    for t in info["tables"])
            grand["summary_rows"] += sum(len(t["summary_rows"]) for t in info["tables"])
            grand["missing"]      += len(info["issues"]["missing"])
            grand["zero"]         += len(info["issues"]["zero"])
            grand["na"]           += len(info["issues"]["na"])
            grand["error"]        += len(info["issues"]["error"])
            grand["fzero"]        += len(info["issues"]["formula_zero"])
            grand["fnone"]        += len(info["issues"]["formula_none"])
        except Exception as e:
            print(f"\n[LỖI] Không đọc được sheet '{sheet}': {e}")

    print("\n" + "─"*78)
    print("  🧾 TỔNG KẾT PHASE 1")
    print("─"*78)
    print(f"  • Tổng số bảng                      : {grand['tables']}")
    print(f"  • Tổng số dòng DỮ LIỆU (đã kiểm tra): {grand['data_rows']}")
    print(f"  • Tổng số dòng TỔNG KẾT (bỏ qua)    : {grand['summary_rows']}")
    print(f"  • Tổng số ô khuyết                  : {grand['missing']}")
    print(f"  • Tổng số ô bằng 0                  : {grand['zero']}")
    print(f"  • Tổng số ô N/A                     : {grand['na']}")
    print(f"  • Tổng số ô lỗi công thức           : {grand['error']}")
    print(f"  • Tổng số ô công thức cho ra 0      : {grand['fzero']}")
    print(f"  • Tổng số ô công thức chưa tính     : {grand['fnone']}")

    # =====================================================
    # LOAD WORKBOOK 1 LẦN DUY NHẤT (dùng cho Phase 2, 3, 4)
    # =====================================================
    try:
        wb_formula = openpyxl.load_workbook(FILE_PATH, data_only=False, rich_text=True)
    except TypeError:
        wb_formula = openpyxl.load_workbook(FILE_PATH, data_only=False)
    wb_data = openpyxl.load_workbook(FILE_PATH, data_only=True)

    # =====================================================
    # PHASE 2
    # =====================================================
    print("\n\n" + "█"*78)
    print("█  PHASE 2 – KIỂM TRA & CHUẨN HÓA HEADER")
    print("█"*78)

    for sheet in sheets:
        try:
            so_header, so_loi, chi_tiet = run_phase2(wb_formula, sheet)
            grand["headers"]       += so_header
            grand["header_errors"] += so_loi

            if so_header == 0:
                continue

            print(f"\n  📄 SHEET: {sheet}")
            print(f"     • Số dòng header nhận diện: {so_header}")
            print(f"     • Số dòng header sai      : {so_loi}")
            for r, ds_loi in chi_tiet:
                chi_tiet_str = ", ".join(
                    f"{chr(65 + i)}={v!r}" for i, v in ds_loi
                )
                print(f"        · Dòng {r}: {chi_tiet_str}")
        except Exception as e:
            print(f"\n[LỖI] Không xử lý được sheet '{sheet}' ở Phase 2: {e}")

    print("\n" + "─"*78)
    print("  🧾 TỔNG KẾT PHASE 2")
    print("─"*78)
    print(f"  • Tổng số dòng header nhận diện : {grand['headers']}")
    print(f"  • Tổng số dòng header bị sai    : {grand['header_errors']}")

    # =====================================================
    # PHASE 3 – FILL MÃ HÀNG
    # =====================================================
    print("\n\n" + "█"*78)
    print("█  PHASE 3 – RESOLVE VLOOKUP & ĐIỀN MÃ HÀNG")
    print("█"*78)

    sheets_all = wb_formula.sheetnames
    sheet_bg = find_sheet(sheets_all, SHEET_BAOGIA)

    # Xác định target_sheets luôn (để Phase 4 dùng dù Phase 3 chạy hay không)
    target_sheets = []
    for s in sheets_all:
        for t in SHEET_TARGETS:
            if s.strip().lower() == t.strip().lower():
                target_sheets.append(s)
                break

    if sheet_bg is None:
        print(f"❌ Không tìm thấy sheet '{SHEET_BAOGIA}' → bỏ qua Phase 3.")
    else:
        required_cols = [COL_MA_HANG, COL_TEN_MAT_HANG, COL_DON_GIA]
        df_bg = read_baogia_auto(FILE_PATH, sheet_bg, required_cols)

        if df_bg is None:
            print("❌ Không đọc được sheet Báo giá → bỏ qua Phase 3.")
        else:
            missing = [c for c in required_cols if c not in df_bg.columns]
            if missing:
                print(f"❌ Sheet '{sheet_bg}' thiếu cột: {missing}")
                print(f"   Cột hiện có: {list(df_bg.columns)}")
            else:
                df_bg = df_bg.dropna(subset=[COL_TEN_MAT_HANG])
                df_bg["_ten"]   = df_bg[COL_TEN_MAT_HANG].apply(norm_text)
                df_bg["_price"] = df_bg[COL_DON_GIA].apply(normalize_price)

                total = len(df_bg)
                if total > 0:
                    nan_count = df_bg["_price"].isna().sum()
                    if nan_count / total > 0.5:
                        print(f"\n⚠️⚠️ Cột '{COL_DON_GIA}' ở '{sheet_bg}' có "
                              f"{nan_count}/{total} ô trống.")

                lookup_exact = {}
                for _, r in df_bg.iterrows():
                    t, p, m = r["_ten"], r["_price"], r[COL_MA_HANG]
                    if t and p is not None and not is_blank(m):
                        lookup_exact.setdefault((t, p), m)

                lookup_name = {}
                for _, r in df_bg.iterrows():
                    t, p, m = r["_ten"], r["_price"], r[COL_MA_HANG]
                    if t and p is not None and not is_blank(m):
                        lookup_name.setdefault(t, []).append((p, m))

                print(f"\n✅ Đã nạp {len(df_bg)} dòng từ '{sheet_bg}'")
                print(f"   - Cặp (Tên + Giá) duy nhất : {len(lookup_exact)}")
                print(f"   - Tên mặt hàng duy nhất    : {len(lookup_name)}")

                if not target_sheets:
                    print(f"⚠️ Không tìm thấy sheet nào trong: {SHEET_TARGETS}")
                else:
                    print(f"\n🎯 Sẽ xử lý {len(target_sheets)} sheet: {target_sheets}")

                    print("\n" + "=" * 65)
                    print("🔄 RESOLVE CÔNG THỨC VLOOKUP (chỉ trên dòng DATA)")
                    print("=" * 65)

                    total_resolved_all = 0
                    for sheet_name in target_sheets:
                        data_rows_set, _ = sheet_classify_map.get(sheet_name, (None, None))
                        n = resolve_vlookups_in_sheet(
                            sheet_name, wb_formula, wb_data, wb_formula, sheet_bg,
                            data_rows_set=data_rows_set
                        )
                        if n > 0:
                            print(f"  ✓ Sheet '{sheet_name}': resolve được {n} ô VLOOKUP")
                        else:
                            print(f"  - Sheet '{sheet_name}': không có VLOOKUP hoặc đã có giá trị sẵn")
                        total_resolved_all += n

                    print(f"\n✅ Tổng cộng đã resolve {total_resolved_all} ô VLOOKUP")

                    for sheet_name in target_sheets:
                        if sheet_name in wb_formula.sheetnames:
                            data_rows_set, summary_rows_set = sheet_classify_map.get(
                                sheet_name, (None, None)
                            )
                            process_sheet(
                                wb_formula, wb_data, wb_formula,
                                sheet_name, lookup_exact, lookup_name,
                                data_rows_set=data_rows_set,
                                summary_rows_set=summary_rows_set,
                            )

    # =====================================================
    # PHASE 4 – XÓA DÒNG DATA CÓ NGUYÊN LIỆU TRỐNG
    # =====================================================
    print("\n\n" + "█"*78)
    print("█  PHASE 4 – XÓA DÒNG DATA CÓ [NGUYÊN LIỆU] TRỐNG")
    print("█"*78)

    if not target_sheets:
        print("  [BỎ QUA] Không có sheet target nào.")
    else:
        for sheet_name in target_sheets:
            if sheet_name not in wb_formula.sheetnames:
                continue

            data_rows_set, _ = sheet_classify_map.get(sheet_name, (None, None))
            if not data_rows_set:
                print(f"\n  📄 SHEET: {sheet_name} → không có classification, bỏ qua.")
                continue

            _, headers, _, _ = _detect_header_cells(wb_formula[sheet_name])
            if not headers or COL_NGUYEN_LIEU not in headers:
                print(f"\n  📄 SHEET: {sheet_name} → không tìm thấy cột "
                      f"'{COL_NGUYEN_LIEU}', bỏ qua.")
                continue

            if DRY_RUN:
                ws = wb_formula[sheet_name]
                col_nl = headers[COL_NGUYEN_LIEU]
                would_delete = [
                    r for r in sorted(data_rows_set)
                    if r <= ws.max_row
                    and is_blank(ws.cell(row=r, column=col_nl).value)
                ]
                print(f"\n  📄 SHEET: {sheet_name}")
                print(f"     (DRY_RUN) Sẽ xóa {len(would_delete)} dòng DATA có NL trống")
                for r in would_delete[:30]:
                    print(f"        · Dòng {r}")
                if len(would_delete) > 30:
                    print(f"        ... và {len(would_delete) - 30} dòng khác")
                grand["deleted_rows"] += len(would_delete)
            else:
                deleted, n_ok, n_skip, n_formula = delete_blank_nl_data_rows(
                    wb_formula, sheet_name, data_rows_set, headers
                )
                print(f"\n  📄 SHEET: {sheet_name}")
                print(f"     • Số dòng đã xóa             : {len(deleted)}")
                print(f"     • Merged cells re-tạo OK     : {n_ok}")
                if n_skip:
                    print(f"     • Merged cells bỏ qua        : {n_skip} "
                          f"(toàn bộ nằm trong vùng xóa)")
                print(f"     • Công thức đã điều chỉnh    : {n_formula} ô")
                for r in deleted[:30]:
                    print(f"        · Dòng {r}")
                if len(deleted) > 30:
                    print(f"        ... và {len(deleted) - 30} dòng khác")
                grand["deleted_rows"]           += len(deleted)
                grand["deleted_merges_ok"]      += n_ok
                grand["deleted_merges_skip"]    += n_skip
                grand["adjusted_formulas"]      += n_formula

        print("\n" + "─"*78)
        print("  🧾 TỔNG KẾT PHASE 4")
        print("─"*78)
        print(f"  • Tổng số dòng đã xóa          : {grand['deleted_rows']}")
        print(f"  • Merged cells re-tạo OK       : {grand['deleted_merges_ok']}")
        print(f"  • Merged cells bỏ qua          : {grand['deleted_merges_skip']}")
        print(f"  • Công thức đã điều chỉnh      : {grand['adjusted_formulas']} ô")

    # =====================================================
    # GHI FILE OUTPUT (DUY NHẤT)
    # =====================================================
    print("\n" + "═"*78)
    if DRY_RUN:
        print("  ℹ️  DRY_RUN = True → KHÔNG ghi file. Đổi thành False để lưu.")
    else:
        wb_formula.save(FILE_OUTPUT)
        print(f"  ✅ Đã lưu file kết quả: {FILE_OUTPUT}")
    print("  ✅ File gốc KHÔNG bị thay đổi.")
    print("═"*78)


if __name__ == "__main__":
    main()